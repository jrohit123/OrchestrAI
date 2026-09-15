"""
excel_engine.py — Plain data-table Excel (.xlsx) generator.

Deliberately NOT LLM-authored like pdf_engine.py — an Excel export is just
the query rows in a sheet, there's no layout to creatively generate, so a
straight openpyxl table is both simpler and faster. No logo either: a
spreadsheet is a data-download for reuse (filtering, pivoting, re-import),
not a branded document to hand someone — the PDF already covers that case.
"""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill(start_color="185FA5", end_color="185FA5", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT  = Font(bold=True, size=14)
SUBTITLE_FONT = Font(italic=True, color="6B7280")

_MAX_COL_WIDTH = 60


def generate_excel(rows: list, title: str, subtitle: str = "") -> bytes:
    """
    Generate a simple .xlsx from `rows` (list of flat dicts — same shape
    query_database/generate_pdf already work with). Column order follows
    the union of keys across all rows, in first-seen order, so a row
    missing a field the others have doesn't shuffle every column.
    Raises ValueError if rows is empty — an empty workbook isn't a useful
    download, better to fail loudly than hand back a blank file.
    """
    if not rows:
        raise ValueError("No rows to export")

    columns: list[str] = []
    seen = set()
    for row in rows:
        for key in row.keys():
            if key not in seen:
                seen.add(key)
                columns.append(key)

    wb = Workbook()
    ws = wb.active
    ws.title = (title or "Sheet1")[:31] or "Sheet1"  # Excel sheet-name cap

    next_row = 1
    if title:
        ws.cell(row=next_row, column=1, value=title).font = TITLE_FONT
        next_row += 1
    if subtitle:
        ws.cell(row=next_row, column=1, value=subtitle).font = SUBTITLE_FONT
        next_row += 1
    if title or subtitle:
        next_row += 1  # blank row before the table

    header_row = next_row
    for col_idx, col_name in enumerate(columns, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=col_name.replace("_", " ").title())
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="left", vertical="center")

    for row_offset, row in enumerate(rows, start=1):
        for col_idx, col_name in enumerate(columns, start=1):
            value = row.get(col_name)
            # openpyxl can't write dicts/lists — stringify anything it
            # can't natively represent rather than raising mid-export.
            if isinstance(value, (dict, list)):
                value = str(value)
            ws.cell(row=header_row + row_offset, column=col_idx, value=value)

    for col_idx, col_name in enumerate(columns, start=1):
        max_len = len(col_name)
        for row in rows:
            cell_val = row.get(col_name)
            if cell_val is not None:
                max_len = max(max_len, len(str(cell_val)))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, _MAX_COL_WIDTH)

    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
